"""The migration target API, and the failure modes of reading template versions."""

from __future__ import annotations

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

from migration_tool import MigrationError, MigrationTarget, migration_target
from migration_tool.outils import VERSION_NR, get_version
from migration_tool.tool import IssuesCollector, _missing_NR, _NR_changes


def _workbook_with_version_cell(value: object | None) -> Workbook:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Introduction"
    if value is not None:
        ws["C1"] = value
    wb.defined_names[VERSION_NR] = DefinedName(
        VERSION_NR, attr_text="Introduction!$C$1"
    )
    return wb


class TestMigrationTarget:
    def test_is_the_bundled_vs_template(self):
        assert migration_target() == MigrationTarget(
            target_template_version="2.0.0",
            target_entry_point="https://xbrl.efrag.org/taxonomy/vs/2026-10-16/vs-all.xsd",
        )

    def test_is_cached(self):
        assert migration_target() is migration_target()


class TestGetVersion:
    def test_returns_the_version(self):
        wb = _workbook_with_version_cell("1.3.0")
        assert get_version(wb, IssuesCollector()) == "1.3.0"

    @pytest.mark.parametrize("value", [None, "", "   "])
    def test_missing_or_blank_version_raises_migration_error(self, value):
        wb = _workbook_with_version_cell(value)
        with pytest.raises(MigrationError, match=VERSION_NR):
            get_version(wb, IssuesCollector())

    def test_missing_named_range_raises_migration_error_not_system_exit(self):
        with pytest.raises(MigrationError):
            get_version(Workbook(), IssuesCollector())


class TestDataFilesAreFoundWithoutRelyingOnCwd:
    @pytest.mark.parametrize("loader", [_missing_NR, _NR_changes])
    def test_loads_from_any_working_directory(self, loader, tmp_path, monkeypatch):
        loader._load.cache_clear()  # otherwise an earlier call hides the problem
        monkeypatch.chdir(tmp_path)
        try:
            assert loader()
        finally:
            loader._load.cache_clear()
