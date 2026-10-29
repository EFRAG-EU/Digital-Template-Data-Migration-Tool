"""The target template's helper cells are found by their row's template label,
so rows inserted in a new template release move them along."""

from __future__ import annotations

from importlib.resources import as_file, files

import pytest
from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.formula import ArrayFormula

from migration_tool.data.validations import VersionCollection
from migration_tool.outils import MigrationError, resolveAnchors
from migration_tool.tool import NEW_TEMPLATE_NAME, _missing_NR, load_workbook_quietly

SHEET = "Social Disclosures"


def _collection(**anchor) -> VersionCollection:
    return VersionCollection.from_dict(
        {"2.0.0": {"Cell": {"sheet": SHEET, "anchor": anchor}}}
    )


@pytest.fixture
def wb() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws["C5"] = "=template_label_training_hours"
    ws["C9"] = ArrayFormula("C9", "=template_label_array_label")
    # a formula that merely mentions a label is not that label's row
    ws["L2"] = "=IF(D5=TRUE, template_label_training_hours, template_label_ok)"
    ws["A12"] = "Other 1"  # sheets without template labels anchor on plain text
    return wb


@pytest.mark.parametrize(
    ("anchor", "expected"),
    [
        ({"label": "template_label_training_hours", "cols": "D"}, "$D$5"),
        (
            {"label": "template_label_training_hours", "cols": "D", "rows": 2},
            "$D$5:$D$6",
        ),
        (
            {"label": "template_label_training_hours", "cols": "D:H", "rows": 2},
            "$D$5:$H$6",
        ),
        (
            {"label": "template_label_training_hours", "cols": "H", "rowOffset": 3},
            "$H$8",
        ),
        ({"label": "template_label_array_label", "cols": "D"}, "$D$9"),
        ({"label": "Other 1", "cols": "B:F", "rows": 2}, "$B$12:$F$13"),
    ],
)
def test_resolves_anchor_to_range(wb, anchor, expected):
    collection = _collection(**anchor)
    resolveAnchors(collection, "2.0.0", wb)
    assert collection.get_ranges("2.0.0") == [expected]


def test_follows_inserted_rows(wb):
    wb[SHEET].insert_rows(1, amount=11)  # openpyxl moves the labels along
    collection = _collection(label="template_label_training_hours", cols="D")
    resolveAnchors(collection, "2.0.0", wb)
    assert collection.get_ranges("2.0.0") == ["$D$16"]


def test_missing_label_is_an_error(wb):
    collection = _collection(label="template_label_not_there", cols="D")
    with pytest.raises(MigrationError, match="template_label_not_there"):
        resolveAnchors(collection, "2.0.0", wb)


def test_label_in_two_rows_is_an_error(wb):
    wb[SHEET]["C7"] = "=template_label_training_hours"
    collection = _collection(label="template_label_training_hours", cols="D")
    with pytest.raises(MigrationError, match="template_label_training_hours"):
        resolveAnchors(collection, "2.0.0", wb)


# --- the shipped data against the shipped template --------------------------


@pytest.fixture(scope="module")
def target() -> Workbook:
    with as_file(files("migration_tool.data").joinpath(NEW_TEMPLATE_NAME)) as path:
        return load_workbook_quietly(path)


@pytest.fixture(scope="module")
def targetRefs(target) -> VersionCollection:
    collection = VersionCollection.from_dict(_missing_NR())
    resolveAnchors(collection, "2.0.0", target)
    return collection


def test_every_target_cell_is_anchored():
    fixed = [
        label
        for label, ref in _missing_NR()["2.0.0"].items()
        if "anchor" not in ref and ref.get("range") is not None
    ]
    assert fixed == []


def test_every_target_cell_is_writable(target, targetRefs):
    """Each range starts on a real cell, not inside a merged block: writing a
    value or comment there would fail (as it did when rows were inserted)."""
    broken = []
    for ref in targetRefs.getRefs("2.0.0"):
        if ref.range is None:
            continue
        left, top, _, _ = range_boundaries(ref.range.replace("$", ""))
        if isinstance(target[ref.sheet].cell(top, left), MergedCell):
            broken.append(f"{ref.label}: {ref.sheet}!{ref.range}")
    assert broken == []
