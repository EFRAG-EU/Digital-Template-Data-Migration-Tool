"""Classified/sensitive-information choices survive migration to the new template.

Before 1.2.0 they are a plain list of disclosure names; from 1.2.0 they are the
Table of Contents checkboxes themselves. Either way each choice must end up as a
ticked checkbox in the new Table of Contents.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from conftest import SAMPLES_DIR
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from migration_tool.tool import _missing_NR, migrate_workbook

TOC = "Table of Contents & Validation"
CHECKBOX_COLUMN = 4
LABEL_COLUMN = 2


def _ticked_labels(workbook) -> list[str]:
    """The (formula) labels of every ticked Table of Contents checkbox."""
    ws = workbook[TOC]
    return [
        str(ws.cell(row, LABEL_COLUMN).value)
        for row in range(1, ws.max_row + 1)
        if ws.cell(row, CHECKBOX_COLUMN).value is True
    ]


def test_list_from_before_1_2_0_ticks_the_matching_checkboxes():
    sample = SAMPLES_DIR / "VSME-Digital-Template-Sample-1.1.1_classifiedinfo.xlsx"
    new_wb, _, _ = migrate_workbook(sample)
    labels = _ticked_labels(new_wb)
    # the sample lists "B4 - Pollution..." and "B9 - Workforce - Health and safety"
    assert any("b4_pollution_of_air_water_and_soil" in lab for lab in labels), labels
    assert any("b9_workforce_health_and_safety" in lab for lab in labels), labels
    assert len(labels) == 2, labels


def _checkbox_rows_in_new_template() -> set[int]:
    rows: set[int] = set()
    for label, ref in _missing_NR()["2.0.0"].items():
        if label.endswith("_ToC"):
            _, top, _, bottom = range_boundaries(ref.split("!")[1].replace("$", ""))
            rows.update(range(top, bottom + 1))
    return rows


@pytest.mark.parametrize("version", ["1.2.0", "1.3.0"])
def test_every_ticked_checkbox_in_1_2_0_and_later_arrives_ticked(
    version: str, tmp_path: Path
):
    """Tick every checkbox in the old Table of Contents and every checkbox the new
    one has must come out ticked: none of the choices are lost on the way."""
    sample = SAMPLES_DIR / f"VSME-Digital-Template-Sample-{version}.xlsx"
    values = load_workbook(sample, data_only=True)[TOC]
    old_checkbox_rows = [
        row
        for row in range(1, values.max_row + 1)
        if values.cell(row, CHECKBOX_COLUMN).value in (True, False)
        and values.cell(row, LABEL_COLUMN).value
    ]
    assert old_checkbox_rows, "sample has no checkboxes?"

    ticked = load_workbook(sample)
    for row in old_checkbox_rows:
        ticked[TOC].cell(row, CHECKBOX_COLUMN).value = True
    ticked_path = tmp_path / "all-ticked.xlsx"
    ticked.save(ticked_path)

    new_wb, _, _ = migrate_workbook(ticked_path)
    new_ws = new_wb[TOC]
    not_ticked = sorted(
        row
        for row in _checkbox_rows_in_new_template()
        if new_ws.cell(row, CHECKBOX_COLUMN).value is not True
    )
    assert not_ticked == []
